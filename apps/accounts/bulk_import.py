"""Spreadsheet parsing and validation for administrator account imports."""

import csv
import hashlib
import io
import re
import secrets
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.utils.dateparse import parse_date
from django.utils import timezone

from .forms import AccountCreateForm
from .models import ACCOUNT_TYPE_CHOICES, User

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 500

FIELD_LABELS = {
    'last_name': 'Last Name',
    'first_name': 'First Name',
    'middle_initial': 'Middle Initial',
    'middle_name': 'Middle Name',
    'suffix': 'Suffix',
    'date_of_birth': 'Date of Birth',
    'sex': 'Sex',
    'email': 'Email Address',
    'contact_number': 'Contact Number',
    'residential_address': 'Residential Address',
    'nationality': 'Nationality',
    'emergency_contact_name': 'Emergency Contact Name',
    'emergency_contact_number': 'Emergency Contact Number',
    'emergency_contact_relationship': 'Emergency Contact Relationship',
    'account_type': 'Account Type',
    'password1': 'Password',
    'password2': 'Password Confirmation',
    'username': 'Username',
}
FIELDS = tuple(FIELD_LABELS)
HEADER_LOOKUP = {
    re.sub(r'[^a-z0-9]', '', name.casefold()): name for name in FIELDS
}
HEADER_LOOKUP.update({
    re.sub(r'[^a-z0-9]', '', label.casefold()): name
    for name, label in FIELD_LABELS.items()
})


def account_import_form(data, creator):
    """Use account form validation without letting a manager elevate the import."""
    form = AccountCreateForm(data=data, creator=creator)
    from accounts.roles import can_create_any_account
    if not can_create_any_account(creator):
        # Match the non-privileged choices exposed by the existing create form,
        # while keeping admin roles unavailable through spreadsheet input.
        form.fields['account_type'].choices = [
            (User.AccountType.STUDENT, 'Student'),
            (User.AccountType.INSTRUCTOR, 'Instructor/Trainer'),
        ]
    return form


def _read_rows(upload):
    if not upload:
        raise ValueError('Choose an Excel (.xlsx) or CSV (.csv) file.')
    if upload.size > MAX_IMPORT_BYTES:
        raise ValueError('The import file must be 5 MB or smaller.')
    suffix = upload.name.rsplit('.', 1)[-1].lower() if '.' in upload.name else ''
    if suffix == 'csv':
        try:
            text = upload.read().decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('The CSV file must use UTF-8 encoding.') from exc
        raw_rows = list(csv.reader(io.StringIO(text)))
    elif suffix == 'xlsx':
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(upload, read_only=True, data_only=True)
            raw_rows = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
        except Exception as exc:
            raise ValueError('Could not read this Excel file. Please use the supplied template.') from exc
    else:
        raise ValueError('Only .xlsx and .csv files are accepted.')

    if not raw_rows:
        raise ValueError('The file has no header row.')
    headers = [HEADER_LOOKUP.get(re.sub(r'[^a-z0-9]', '', str(value or '').strip().casefold()))
               for value in raw_rows[0]]
    if not any(headers):
        raise ValueError('The first row must contain the template column headings.')
    if len([field for field in headers if field]) != len(set(field for field in headers if field)):
        raise ValueError('The file contains duplicate column headings.')
    rows = []
    for row_number, cells in enumerate(raw_rows[1:], start=2):
        values = dict(zip(headers, cells))
        if not any(value not in (None, '') for value in values.values()):
            continue
        rows.append((row_number, values))
    if not rows:
        raise ValueError('The file has no account records to import.')
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValueError(f'Import up to {MAX_IMPORT_ROWS} accounts per file.')
    return rows


def _choice_code(value, choices):
    """Accept either a stored choice code or the label shown in the UI/template."""
    normalized = re.sub(r'[^a-z0-9]', '', str(value or '').casefold())
    for code, label in choices:
        if normalized in {
            re.sub(r'[^a-z0-9]', '', str(code).casefold()),
            re.sub(r'[^a-z0-9]', '', str(label).casefold()),
        }:
            return code
    return value


def _form_values(values):
    data = {}
    for name in FIELDS:
        value = values.get(name)
        if value is None:
            value = ''
        if name == 'date_of_birth' and value:
            if hasattr(value, 'strftime'):
                value = value.strftime('%Y-%m-%d')
            else:
                value = str(value).strip()
                parsed = parse_date(value)
                if parsed:
                    value = parsed.isoformat()
        else:
            value = str(value).strip()
        if name == 'sex' and value:
            value = _choice_code(value, User._meta.get_field('sex').choices)
        elif name == 'nationality' and value:
            value = _choice_code(value, User._meta.get_field('nationality').choices)
        elif name == 'account_type' and value:
            value = _choice_code(value, ACCOUNT_TYPE_CHOICES)
        elif name in {'contact_number', 'emergency_contact_number'}:
            # Excel treats phone numbers as numeric cells and drops the leading
            # zero from Philippine mobile numbers. Restore it before validation.
            digits = re.sub(r'\D', '', value)
            if len(digits) == 10 and digits.startswith('9'):
                value = '0' + digits
        data[name] = value
    return data


def validate_upload(upload, creator):
    raw_rows = _read_rows(upload)
    parsed = []
    username_rows = {}
    email_rows = {}
    User = get_user_model()
    existing_usernames = set()
    existing_emails = set()
    for existing_username, existing_email in User.objects.values_list('username', 'email'):
        existing_usernames.add(existing_username.casefold())
        if existing_email:
            existing_emails.add(existing_email.casefold())
    reserved_usernames = set(existing_usernames)
    reserved_usernames.update(
        str(values.get('username') or '').strip().casefold()
        for _, values in raw_rows if values.get('username')
    )
    for row_number, values in raw_rows:
        data = _form_values(values)
        username_generated = not bool(data['username'])
        password_generated = not data['password1'] and not data['password2']
        if not data['username'] and data['email']:
            base = re.sub(r'[^a-zA-Z0-9_]', '', data['email'].split('@')[0])[:140] or 'user'
            username = base
            suffix = 1
            while username.casefold() in reserved_usernames:
                username = f'{base[:145 - len(str(suffix))]}{suffix}'
                suffix += 1
            data['username'] = username
            reserved_usernames.add(username.casefold())
        if password_generated:
            generated_password = secrets.token_urlsafe(24)
            data['password1'] = generated_password
            data['password2'] = generated_password
        username = data['username'].casefold()
        email = data['email'].casefold()
        if username:
            username_rows.setdefault(username, []).append(row_number)
        if email:
            email_rows.setdefault(email, []).append(row_number)
        parsed.append((row_number, data, username_generated, password_generated))

    result = []
    valid_forms = {}
    for row_number, data, username_generated, password_generated in parsed:
        form = account_import_form(data=data, creator=creator)
        form.is_valid()
        errors = [f'{field}: {message}' for field, messages in form.errors.items()
                  for message in messages]
        username = data['username'].casefold()
        email = data['email'].casefold()
        if len(username_rows.get(username, [])) > 1:
            errors.append('Username is duplicated within this file.')
        if len(email_rows.get(email, [])) > 1:
            errors.append('Email is duplicated within this file.')
        if username and username in existing_usernames:
            errors.append('Username is already in use.')
        if email and email in existing_emails:
            errors.append('Email address is already in use.')
        valid_forms[row_number] = form
        result.append({
            'row': row_number,
            'first_name': data['first_name'],
            'last_name': data['last_name'],
            'username': data['username'],
            'email': data['email'],
            'account_type': data['account_type'],
            'username_generated': username_generated,
            'password_setup_required': password_generated,
            'valid': not errors,
            'errors': errors,
            'status': 'Ready' if not errors else 'Needs correction',
        })
    return result, valid_forms


def create_valid_accounts(upload, creator):
    results, forms = validate_upload(upload, creator)
    by_row = {row['row']: row for row in results}
    created = 0
    setup_email_failed = 0
    for row_number, form in forms.items():
        result = by_row[row_number]
        if not result['valid']:
            continue
        setup_token_id = None
        try:
            with transaction.atomic():
                user = form.save()
                if result['password_setup_required']:
                    setup_token_id = _queue_password_setup(user)
        except (IntegrityError, ValidationError):
            result['valid'] = False
            result['status'] = 'Needs correction'
            result['errors'].append('Username or email is already in use.')
            continue
        setup_email_sent = True
        if setup_token_id:
            from notifications.models import PasswordResetToken
            setup_email_sent = PasswordResetToken.objects.filter(
                pk=setup_token_id, email_status='SENT').exists()
            if not setup_email_sent:
                result['status'] = 'Created; setup email failed'
                result['errors'].append(
                    'Password setup email could not be sent; check outbound email configuration.')
                setup_email_failed += 1
        result['status'] = 'Created' if setup_email_sent else 'Created; setup email failed'
        created += 1
    return results, created, setup_email_failed


def _queue_password_setup(user):
    """Queue the existing one-use password reset link for a generated password."""
    from notifications.models import PasswordResetToken

    raw = secrets.token_urlsafe(32)
    now = timezone.now()
    PasswordResetToken.objects.filter(user=user, used_at__isnull=True).update(used_at=now)
    reset = PasswordResetToken.objects.create(
        user=user,
        token_hash=hashlib.sha256(raw.encode()).hexdigest(),
        expires_at=now + timedelta(minutes=30),
    )
    from notifications.email_service import deliver_password_reset
    transaction.on_commit(lambda pk=reset.pk, token=raw: deliver_password_reset(pk, token))
    return reset.pk


def template_headers():
    return [FIELD_LABELS[field] for field in FIELDS]
